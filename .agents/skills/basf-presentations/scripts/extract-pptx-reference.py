from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


EMU_PER_INCH = 914400


def emu_to_inch(value: int) -> float:
    return round(value / EMU_PER_INCH, 3)


def shape_type_name(shape) -> str:
    try:
        return MSO_SHAPE_TYPE(shape.shape_type).name
    except Exception:
        return str(shape.shape_type)


def collect_text(shape) -> tuple[list[str], Counter, Counter]:
    samples: list[str] = []
    fonts: Counter = Counter()
    sizes: Counter = Counter()

    if not shape.has_text_frame:
        return samples, fonts, sizes

    paragraphs = [paragraph.text for paragraph in shape.text_frame.paragraphs]
    text = "\n".join(paragraphs).strip()
    if text:
        samples.append(text)

    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                fonts[run.font.name] += 1
            if run.font.size:
                sizes[round(run.font.size.pt, 2)] += 1

    return samples, fonts, sizes


def export_picture(shape, slide_dir: Path, slide_number: int, shape_index: int) -> dict:
    image_info: dict[str, object] = {}
    try:
        image = shape.image
    except Exception as exc:
        image_info["embedded"] = False
        image_info["error"] = str(exc)
        return image_info

    image_name = f"slide-{slide_number:02d}-shape-{shape_index:02d}.{image.ext}"
    image_path = slide_dir / image_name
    image_path.write_bytes(image.blob)
    image_info["embedded"] = True
    image_info["path"] = str(image_path)
    image_info["ext"] = image.ext
    image_info["size_bytes"] = len(image.blob)
    image_info["dpi"] = image.dpi
    return image_info


def export_table(shape, slide_dir: Path, slide_number: int, shape_index: int) -> dict:
    rows = []
    for row in shape.table.rows:
        rows.append([cell.text_frame.text.strip() for cell in row.cells])

    csv_path = slide_dir / f"slide-{slide_number:02d}-shape-{shape_index:02d}-table.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)

    return {
        "rows": len(shape.table.rows),
        "cols": len(shape.table.columns),
        "path": str(csv_path),
        "preview": rows[:6],
    }


def export_chart(shape, slide_dir: Path, slide_number: int, shape_index: int) -> dict:
    chart = shape.chart
    categories = []
    workbook = chart.plots[0].categories
    for category in workbook:
        categories.append(str(category.label))

    series = []
    for item in chart.series:
        series.append(
            {
                "name": item.name,
                "values": list(item.values),
            }
        )

    chart_data = {
        "type": str(chart.chart_type),
        "series_count": len(chart.series),
        "categories": categories,
        "series": series,
    }
    chart_path = slide_dir / f"slide-{slide_number:02d}-shape-{shape_index:02d}-chart.json"
    chart_path.write_text(json.dumps(chart_data, indent=2), encoding="utf-8")
    chart_data["path"] = str(chart_path)
    return chart_data


def analyze_slide(slide, slide_number: int, output_dir: Path) -> dict:
    slide_dir = output_dir / f"slide-{slide_number:02d}"
    slide_dir.mkdir(parents=True, exist_ok=True)

    shape_counts: Counter = Counter()
    fonts: Counter = Counter()
    sizes: Counter = Counter()
    text_samples: list[str] = []
    details: list[dict] = []
    pictures = 0
    tables = 0
    charts = 0

    for shape_index, shape in enumerate(slide.shapes, start=1):
        shape_counts[shape_type_name(shape)] += 1
        samples, shape_fonts, shape_sizes = collect_text(shape)
        fonts.update(shape_fonts)
        sizes.update(shape_sizes)
        text_samples.extend(samples[:1])

        detail = {
            "index": shape_index,
            "name": getattr(shape, "name", ""),
            "type": shape_type_name(shape),
            "placeholder": bool(getattr(shape, "is_placeholder", False)),
            "box_in": {
                "left": emu_to_inch(getattr(shape, "left", 0)),
                "top": emu_to_inch(getattr(shape, "top", 0)),
                "width": emu_to_inch(getattr(shape, "width", 0)),
                "height": emu_to_inch(getattr(shape, "height", 0)),
            },
        }

        if samples:
            detail["text"] = samples[:1]
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures += 1
            detail["image"] = export_picture(shape, slide_dir, slide_number, shape_index)
        if shape.has_table:
            tables += 1
            detail["table"] = export_table(shape, slide_dir, slide_number, shape_index)
        if getattr(shape, "has_chart", False):
            charts += 1
            detail["chart"] = export_chart(shape, slide_dir, slide_number, shape_index)

        details.append(detail)

    return {
        "slide_number": slide_number,
        "layout_name": slide.slide_layout.name,
        "shape_count": len(slide.shapes),
        "shape_counts": dict(shape_counts),
        "pictures": pictures,
        "tables": tables,
        "charts": charts,
        "top_fonts": fonts.most_common(10),
        "top_font_sizes_pt": sizes.most_common(12),
        "text_samples": text_samples[:10],
        "shape_details": details,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx_path", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument(
        "--slides",
        type=str,
        default="",
        help="Comma-separated 1-based slide numbers to export. Default exports all slides.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    prs = Presentation(args.pptx_path)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    selected = None
    if args.slides:
        selected = {int(value.strip()) for value in args.slides.split(",") if value.strip()}

    slides = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        if selected and slide_number not in selected:
            continue
        slides.append(analyze_slide(slide, slide_number, args.output_dir))

    summary = {
        "file": str(args.pptx_path),
        "slide_size": {
            "width_emu": prs.slide_width,
            "height_emu": prs.slide_height,
            "width_in": emu_to_inch(prs.slide_width),
            "height_in": emu_to_inch(prs.slide_height),
        },
        "slide_count": len(prs.slides),
        "exported_slide_count": len(slides),
        "slides": slides,
    }
    summary_path = args.output_dir / "summary.json"
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(summary_path)


if __name__ == "__main__":
    main()
